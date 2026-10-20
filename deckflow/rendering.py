"""共用排版工具；CJK 逐字符估宽，图片与文字使用独立内容框。"""
from dataclasses import dataclass
from pathlib import Path
import hashlib
import math
import unicodedata

from PIL import Image, ImageDraw, ImageOps
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from .content import WorkflowError
from .strings import tr


@dataclass(frozen=True)
class Rect:
    x: float
    y: float
    w: float
    h: float


def char_width(char, size):
    if unicodedata.combining(char):
        return 0
    if unicodedata.east_asian_width(char) in ("F", "W"):
        return size
    if char in "il.,:;'!| ":
        return size * 0.34
    if char in "MW@%":
        return size * 0.9
    return size * 0.62


def line_count(text, width_pt, size):
    # 对中文逐字符贪心换行；换行符与空段落都计入高度。
    if width_pt <= 0:
        return math.inf
    lines = 0
    for paragraph in text.split("\n"):
        used = 0
        lines += 1
        for char in paragraph:
            width = char_width(char, size)
            if used + width > width_pt and used > 0:
                lines += 1
                used = 0
            used += width
    return lines


def text_height(text, width_in, size, leading=1.32):
    return line_count(text, width_in * 72, size) * size * leading / 72


def fit_size(text, rect, requested, minimum, page=None, field="text"):
    size = requested
    while size >= minimum:
        if text_height(text, rect.w, size) <= rect.h:
            return size
        size -= 0.5
    raise WorkflowError("V9", tr("text_too_long", needed=text_height(text, rect.w, minimum),
                                 available=rect.h), page, field)


def set_role(shape, role):
    # 不依赖每页重复的形状名；只访问对应 XML 的直接子节点。
    for nv in shape.element:
        if nv.tag.rsplit("}", 1)[-1].startswith("nv"):
            for node in nv:
                if node.tag.rsplit("}", 1)[-1] == "cNvPr":
                    node.set("descr", "deckflow:" + role)
                    return


def get_role(shape):
    for nv in shape.element:
        if nv.tag.rsplit("}", 1)[-1].startswith("nv"):
            for node in nv:
                if node.tag.rsplit("}", 1)[-1] == "cNvPr":
                    return node.get("descr", "").removeprefix("deckflow:")
    return ""


def style_font(font, size, color, family, bold=False):
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = RGBColor.from_string(color)
    # 字体元素由 python-pptx 按 OOXML 序列插入，不能放在填充之前。
    font.name = family
    props = font._rPr
    props.set("lang", "zh-CN")
    ea = next((node for node in props if node.tag.endswith("}ea")), None)
    if ea is None:
        ea = OxmlElement("a:ea")
        props.insert_element_before(ea, "a:cs", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl")
    ea.set("typeface", family)


def fill_frame(frame, text, size, color, family, bold=False):
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for attr in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(frame, attr, 0)
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.line_spacing = 1.12
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        style_font(paragraph.font, size, color, family, bold)
        run = paragraph.add_run()
        run.text = line
        style_font(run.font, size, color, family, bold)


class Renderer:
    def __init__(self, slide, course, page, theme, total):
        self.slide, self.course, self.page = slide, course, page
        self.theme, self.total = theme, total
        self.palette, self.fonts = theme["palette"], theme["fonts"]
        self.geo, self.type = theme["geometry"], theme["type_scale"]
        self.dark = page.layout in {"cover", "dark", "section"}
        self.ink = "offwhite" if self.dark else "ink"
        self.soft = "accent_mid" if self.dark else "ink_soft"

    def text(self, rect, text, size=None, color=None, bold=False, role="content", minimum=None):
        if not text:
            return None
        size = size or self.type["body"]
        fields = list(dict.fromkeys(key for key, value in self.page.entries
                      if key not in {"layout", "time", "image", "image2"} and value in text))
        actual = fit_size(text, rect, size, minimum or min(size, self.type["min_body"]),
                          self.page.number, "/".join(fields) or role)
        box = self.slide.shapes.add_textbox(Inches(rect.x), Inches(rect.y), Inches(rect.w), Inches(rect.h))
        family = "head" if bold and size >= self.type["h1"] else "cjk"
        if role == "chrome" and text.isascii():
            family = "mono"
        fill_frame(box.text_frame, text, actual, self.palette[color or self.ink],
                   self.fonts[family], bold)
        set_role(box, role)
        return box

    def panel(self, rect, color="paper_alt", line=None, rounded=False):
        shape = self.slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
            Inches(rect.x), Inches(rect.y), Inches(rect.w), Inches(rect.h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(self.palette[color])
        shape.line.fill.background()
        shape.shadow.inherit = False
        if line:
            shape.line.color.rgb = RGBColor.from_string(self.palette[line])
        set_role(shape, "decor")
        return shape

    def title(self):
        m, width = self.geo["margin"], self.geo["slide_w"] - 2 * self.geo["margin"]
        self.text(Rect(m, .87, width, .54), self.page.get("title"), self.type["h1"], bold=True)
        self.text(Rect(m, 1.49, width, .29), self.page.get("subtitle"), self.type["small"], color=self.soft)

    def content_area(self, key=True):
        self.title()
        top, bottom, margin = self.geo["content_top"], self.geo["content_bottom"], self.geo["margin"]
        width = self.geo["slide_w"] - 2 * margin
        if key and self.page.get("key"):
            self.panel(Rect(margin, top, width, .72), "accent" if not self.dark else "ink")
            self.text(Rect(margin + .18, top + .13, width - .36, .46), self.page.get("key"),
                      self.type["body"], "white", True)
            top += .94
        return Rect(margin, top, width, bottom - top)

    def side_images(self, area, width=3.65):
        images = self.page.images
        if not images:
            return area
        if len(images) > 3:
            raise WorkflowError("INPUT", tr("too_many_images"), self.page.number, "image")
        gap = .24
        x = area.x + area.w - width
        height = (area.h - gap * (len(images) - 1)) / len(images)
        for index, (name, caption) in enumerate(images):
            self.picture_with_caption(Rect(x, area.y + index * (height + gap), width, height), name, caption)
        return Rect(area.x, area.y, area.w - width - .38, area.h)

    def image(self, rect, name, cover=False):
        path = self.course.image_path(name, self.page.number)
        cache = self.course.output / ".cache"
        cache.mkdir(parents=True, exist_ok=True)
        try:
            with Image.open(path) as source:
                image = ImageOps.exif_transpose(source).convert("RGB")
            if cover:
                image = ImageOps.fit(image, (2880, 1620), method=Image.Resampling.LANCZOS)
                # 压暗烘进像素；不使用半透明 PPT 遮罩。
                shade = Image.new("L", image.size)
                draw = ImageDraw.Draw(shade)
                for x in range(image.width):
                    alpha = int(170 - 95 * x / image.width)
                    draw.line((x, 0, x, image.height), fill=alpha)
                image = Image.composite(Image.new("RGB", image.size, "#101B16"), image, shade)
            digest = hashlib.sha256(image.tobytes()).hexdigest()[:20]
            normalized = cache / (digest + ".png")
            if not normalized.exists():
                image.save(normalized)
            if cover:
                target = rect
            else:
                scale = min(rect.w / image.width, rect.h / image.height)
                w, h = image.width * scale, image.height * scale
                target = Rect(rect.x + (rect.w - w) / 2, rect.y + (rect.h - h) / 2, w, h)
            shape = self.slide.shapes.add_picture(str(normalized), Inches(target.x), Inches(target.y),
                                                  width=Inches(target.w), height=Inches(target.h))
            set_role(shape, "background" if cover else "content")
            return shape
        except (OSError, ValueError) as exc:
            raise WorkflowError("V8", tr("image_error", path=path, error=exc), self.page.number, "image") from exc

    def picture_with_caption(self, rect, name, caption):
        cap_height = .48 if caption else 0
        image_rect = Rect(rect.x, rect.y, rect.w, rect.h - cap_height - (.10 if caption else 0))
        self.panel(image_rect)
        self.image(image_rect, name)
        if caption:
            self.text(Rect(rect.x, rect.y + rect.h - cap_height, rect.w, cap_height), caption,
                      self.type["caption"], self.soft, minimum=9)

    def rows(self, area, items, numbered=False):
        if not items:
            return
        gap = .15
        height = (area.h - gap * (len(items) - 1)) / len(items)
        for index, item in enumerate(items):
            y = area.y + index * (height + gap)
            if numbered:
                self.text(Rect(area.x, y + .02, .55, min(.45, height)), f"{index + 1:02}",
                          18, "accent_mid" if self.dark else "accent", True)
                x, width = area.x + .75, area.w - .75
            else:
                self.panel(Rect(area.x, y + .08, .055, min(.3, height)), "accent")
                x, width = area.x + .22, area.w - .22
            self.text(Rect(x, y, width, height), item, self.type["body"])

    def chrome(self):
        m, width, height = self.geo["margin"], self.geo["slide_w"], self.geo["slide_h"]
        color = self.soft
        self.text(Rect(m, .27, width * .38, .24), self.course.meta["org"], 9, color, role="chrome")
        self.text(Rect(width * .48, .27, width * .27, .24), self.course.meta["docid"], 9, color, role="chrome")
        self.text(Rect(width - m - .82, .27, .82, .24), f"{self.page.number:02} / {self.total:02}",
                  9, color, role="chrome")
        self.panel(Rect(m, .65, width - 2 * m, .014), "accent_mid")
        self.panel(Rect(m, height - .49, width - 2 * m, .012), "accent_mid")
        show_timing = self.theme.get("show_timing", True)
        footer_width = width - 2 * m - (2.25 if show_timing else 0)
        self.text(Rect(m, height - .36, footer_width, .25), self.page.get("footer"),
                  8, color, role="chrome", minimum=6)
        if show_timing:
            self.text(Rect(width - m - 2.10, height - .36, 2.10, .25),
                      tr("time_label", minutes=self.page.seconds // 60, seconds=self.page.seconds % 60),
                      8, color, role="chrome")


def extract_slide_text(slide):
    result = []
    for shape in slide.shapes:
        if get_role(shape) != "content":
            continue
        if shape.has_text_frame and shape.text.strip():
            result.append(shape.text)
        if shape.has_table:
            for row in shape.table.rows:
                result.append(" | ".join(cell.text for cell in row.cells))
    return result
