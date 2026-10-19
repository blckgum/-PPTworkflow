"""布局注册表；布局不处理课程专属内容与页眉页脚。"""
import math
import re
from pptx.dml.color import RGBColor
from pptx.util import Inches

from .rendering import Rect, fill_frame, fit_size, set_role
from .strings import tr

LAYOUTS = {}


def register(name):
    def decorator(function):
        LAYOUTS[name] = function
        return function
    return decorator


@register("cover")
def cover(r):
    g, p = r.geo, r.page
    r.image(Rect(0, 0, g["slide_w"], g["slide_h"]), p.images[0][0], cover=True)
    m, width = g["margin"], g["slide_w"] - 2 * g["margin"]
    r.text(Rect(m, 1.25, width, 1.0), p.get("title"), r.type["cover"], "white", True)
    r.text(Rect(m, 2.42, width, .58), p.get("subtitle"), 26, "offwhite")
    r.text(Rect(m, 3.13, width, .44), p.get("subtitle2"), 16, "offwhite")
    r.text(Rect(m, 4.12, width * .88, .85), p.get("question"), 23, "white", True)
    r.text(Rect(m, 5.22, width * .88, .7), p.get("key"), 18, "offwhite")
    r.text(Rect(m, 6.17, width, .40), p.images[0][1], 9, "offwhite", minimum=8)


@register("statement")
def statement(r):
    p, g = r.page, r.geo
    m = g["margin"]
    area = Rect(m, 1.9, g["slide_w"] - 2 * m, 3.6)
    area = r.side_images(area)
    r.text(area, p.get("title"), 42, bold=True)
    r.text(Rect(m, 5.78, g["slide_w"] - 2 * m, .5), p.get("subtitle"), 20, r.soft)
    if p.get("key"):
        r.text(Rect(m, 6.36, g["slide_w"] - 2 * m, .38), p.get("key"), 14, "accent", True)


@register("points")
def points(r):
    r.rows(r.side_images(r.content_area()), r.page.all("bullet") + r.page.all("body"))


@register("blocks")
def blocks(r):
    area = r.side_images(r.content_area())
    heads, bodies = r.page.all("head"), r.page.all("body")
    gap = .18
    height = (area.h - gap * (len(heads) - 1)) / len(heads)
    for index, (head, body) in enumerate(zip(heads, bodies)):
        y = area.y + index * (height + gap)
        r.panel(Rect(area.x, y, area.w, height))
        # 标题与正文共用文本框，换行后仍按内容实际高度适配。
        box = r.text(Rect(area.x + .20, y + .12, area.w - .4, height - .24),
                     head + "\n" + body, 19)
        for run in box.text_frame.paragraphs[0].runs:
            run.font.bold = True


@register("columns")
def columns(r):
    area = r.side_images(r.content_area())
    values = r.page.all("column")
    gap = .25
    width = (area.w - gap * (len(values) - 1)) / len(values)
    for index, value in enumerate(values):
        title, body = value.split("|", 1)
        x = area.x + index * (width + gap)
        r.panel(Rect(x, area.y, width, area.h))
        r.text(Rect(x + .20, area.y + .18, width - .4, .70), title, 21, "accent", True)
        r.text(Rect(x + .20, area.y + 1.05, width - .4, area.h - 1.25), body, 18)


@register("table")
def table(r):
    area = r.side_images(r.content_area())
    rows = [[cell.strip() for cell in value.split("|")] for value in r.page.all("row")]
    shape = r.slide.shapes.add_table(len(rows), len(rows[0]), Inches(area.x), Inches(area.y),
                                     Inches(area.w), Inches(area.h))
    set_role(shape, "content")
    table_obj = shape.table
    table_obj.first_row = False
    table_obj.horz_banding = False
    for row_index, values in enumerate(rows):
        table_obj.rows[row_index].height = Inches(area.h / len(rows))
        for column, value in enumerate(values):
            cell = table_obj.cell(row_index, column)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(r.palette[
                "accent" if row_index == 0 else "paper_alt" if row_index % 2 else "paper"])
            cell.margin_left = cell.margin_right = Inches(.12)
            cell.margin_top = cell.margin_bottom = Inches(.10)
            rect = Rect(0, 0, area.w / len(values) - .24, area.h / len(rows) - .20)
            size = fit_size(value, rect, 17, 11, r.page.number, "row")
            fill_frame(cell.text_frame, value, size, r.palette["white" if row_index == 0 else "ink"],
                       r.fonts["cjk"], row_index == 0)
            cell.margin_left = cell.margin_right = Inches(.12)
            cell.margin_top = cell.margin_bottom = Inches(.10)


@register("kv")
def kv(r):
    area = r.side_images(r.content_area())
    values = r.page.all("kv")
    height = area.h / len(values)
    for index, value in enumerate(values):
        key, body = value.split("=", 1)
        y = area.y + index * height
        r.text(Rect(area.x, y + .08, area.w * .27, height - .16), key, 18, "accent", True)
        r.text(Rect(area.x + area.w * .30, y + .08, area.w * .70, height - .16), body, 18)
        r.panel(Rect(area.x, y + height - .02, area.w, .012), "accent_mid")


@register("chain")
def chain(r):
    r.rows(r.side_images(r.content_area()), r.page.all("step"), numbered=True)


@register("flow")
def flow(r):
    area = r.side_images(r.content_area())
    values = r.page.all("step")
    # 长流程自动分行，避免横向压缩成不可读的窄条。
    cols = min(4, len(values))
    row_count = math.ceil(len(values) / cols)
    gap_x, gap_y = .40, .35
    width = (area.w - gap_x * (cols - 1)) / cols
    height = (area.h - gap_y * (row_count - 1)) / row_count
    for index, value in enumerate(values):
        row, col = divmod(index, cols)
        x, y = area.x + col * (width + gap_x), area.y + row * (height + gap_y)
        r.panel(Rect(x, y, width, height))
        r.text(Rect(x + .16, y + .15, width - .32, .40), f"{index + 1:02}", 18, "accent", True)
        r.text(Rect(x + .16, y + .70, width - .32, height - .86), value, 18)
        if col < cols - 1 and index < len(values) - 1:
            r.text(Rect(x + width + .06, y + height / 2 - .16, .28, .35), "\u2192", 18, "accent", role="decor")


@register("diagram")
def diagram(r):
    area = r.content_area()
    images = r.page.images
    text_width = 3.6 if r.page.all("body") else 0
    picture_width = area.w - (text_width + .35 if text_width else 0)
    gap = .20
    height = (area.h - gap * (len(images) - 1)) / len(images)
    for index, (name, caption) in enumerate(images):
        r.picture_with_caption(Rect(area.x, area.y + index * (height + gap), picture_width, height), name, caption)
    if text_width:
        r.rows(Rect(area.x + picture_width + .35, area.y, text_width, area.h), r.page.all("body"))


@register("grid")
def grid(r):
    area = r.side_images(r.content_area())
    values = r.page.all("card")
    rows = math.ceil(len(values) / 2)
    gap = .20
    width, height = (area.w - gap) / 2, (area.h - gap * (rows - 1)) / rows
    for index, value in enumerate(values):
        title, body = value.split("|", 1)
        row, col = divmod(index, 2)
        x, y = area.x + col * (width + gap), area.y + row * (height + gap)
        r.panel(Rect(x, y, width, height))
        r.text(Rect(x + .18, y + .15, width - .36, .48), title, 18, "accent", True)
        r.text(Rect(x + .18, y + .75, width - .36, height - .9), body, 17)


@register("dark")
def dark(r):
    area = r.side_images(r.content_area())
    r.text(Rect(area.x, area.y, area.w, 1.15), r.page.get("statement"), 31, "white", True)
    r.rows(Rect(area.x, area.y + 1.5, area.w, area.h - 1.5), r.page.all("reason"))


@register("summary")
def summary(r):
    items = [value if re.match(r"\d+\s", value) else f"{index + 1:02}  {value}"
             for index, value in enumerate(r.page.all("num"))]
    r.rows(r.side_images(r.content_area()), items)


@register("closing")
def closing(r):
    g, p = r.geo, r.page
    m = g["margin"]
    area = r.side_images(Rect(m, 1.65, g["slide_w"] - 2 * m, 3.5))
    r.text(Rect(area.x, area.y, area.w, 1.25), p.get("title"), 38, bold=True)
    r.text(Rect(area.x, area.y + 1.65, area.w, 1.25), p.get("title2"), 38, "accent", True)
    r.text(Rect(m, 5.53, g["slide_w"] - 2 * m, .55), p.get("subtitle"), 19, r.soft)
    r.text(Rect(m, 6.20, g["slide_w"] - 2 * m, .45), p.get("key"), 16, "accent")


@register("sources")
def sources(r):
    r.rows(r.side_images(r.content_area()), r.page.all("source"), numbered=True)


@register("qa")
def qa(r):
    area = r.side_images(r.content_area())
    r.rows(area, r.page.all("body"))


@register("section")
def section(r):
    statement(r)
