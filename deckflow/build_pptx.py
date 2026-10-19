"""PPT 生成：布局与统一装饰分开处理。"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

from .layouts import LAYOUTS
from .rendering import Renderer, Rect


def build_pptx(course, theme):
    course.output.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    geo = theme["geometry"]
    presentation.slide_width = Inches(geo["slide_w"])
    presentation.slide_height = Inches(geo["slide_h"])
    presentation.core_properties.title = course.meta["title"]
    presentation.core_properties.subject = course.meta["subtitle"]
    presentation.core_properties.author = course.meta["org"]
    presentation.core_properties.identifier = course.meta["docid"]
    presentation.core_properties.keywords = course.meta["deckname"]
    for page in course.pages:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(
            theme["palette"]["dark_bg" if page.layout in {"dark", "section"} else "paper"])
        renderer = Renderer(slide, course, page, theme, len(course.pages))
        LAYOUTS[page.layout](renderer)
        if page.get("note"):
            # 主讲提示单列底部区域；有提示时由布局预留空间。
            from .strings import tr
            renderer.text(Rect(geo["margin"], 6.82, geo["slide_w"] - 2 * geo["margin"], .17),
                          tr("note_label") + " / " + page.get("note"), 8, "emphasis", minimum=7)
        renderer.chrome()
        slide.notes_slide.notes_text_frame.text = page.notes
    presentation.save(course.pptx_path)
    return course.pptx_path
