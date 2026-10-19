"""输入预检与成品 V1–V10 校验，失败不能被警告掩盖。"""
from dataclasses import dataclass
from pathlib import Path
import re

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .build_docx import read_document_units
from .rendering import extract_slide_text, get_role, text_height
from .source_lint import lint_sources
from .strings import tr

EMU = 914400


@dataclass
class Result:
    code: str
    status: str
    message: str


def report(results):
    print(tr("result_header"))
    for result in results:
        print(f"[{tr(result.status)}] {result.code:<6} {result.message}")


def passed(results):
    return all(result.status != "fail" for result in results)


def preflight(course, engine_root):
    results = []
    total = sum(page.seconds for page in course.pages)
    minimum, maximum = (float(course.meta[key]) for key in ("time_target_min", "time_target_max"))
    results.append(Result("V2", "pass" if minimum * 60 <= total <= maximum * 60 else "fail",
                          tr("v2", minutes=total // 60, seconds=total % 60,
                             minimum=f"{minimum:g}", maximum=f"{maximum:g}")))
    missing = [page for page in course.pages if not page.notes.strip()]
    for page in missing:
        results.append(Result("V3", "fail", tr("page", page=page.number) + " " + tr(
            "v3_missing", path=course.root / "notes" / f"{page.number:02}.txt")))
    if not missing:
        results.append(Result("V3", "pass", tr("v3_source_ok", count=len(course.pages))))
    lint = lint_sources(Path(engine_root))
    results.extend(Result("V7", "fail", message) for message in lint)
    if not lint:
        results.append(Result("V7", "pass", tr("v7_ok")))
    count, invalid = 0, False
    for page in course.pages:
        for name, _ in page.images:
            count += 1
            path = course.image_path(name, page.number)
            if not path.is_file():
                invalid = True
                results.append(Result("V8", "fail", tr("page", page=page.number) + " " +
                                      tr("field", field="image") + " " + tr("v8_missing", path=path)))
                continue
            try:
                with Image.open(path) as image:
                    width, height = image.size
                    image.verify()
                mp = width * height / 1_000_000
                if page.number == 1 and mp < 2.07:
                    results.append(Result("V5", "fail", tr("v5_bad") + f" ({path}: {mp:.2f} MP)"))
                elif page.number != 1 and mp < .8:
                    results.append(Result("V8", "warn", tr("image_small", path=path, mp=mp)))
            except (OSError, ValueError) as exc:
                invalid = True
                results.append(Result("V8", "fail", tr("image_error", path=path, error=exc)))
    if not invalid:
        results.append(Result("V8", "pass", tr("v8_ok", count=count)))
    return results


def check_geometry(presentation):
    results = []
    for page, slide in enumerate(presentation.slides, 1):
        for index, shape in enumerate(slide.shapes, 1):
            # 使用 API 读取直接变换，不搜索嵌套几何中的 a:xfrm。
            if (shape.left < -.01 * EMU or shape.top < -.01 * EMU or
                    shape.left + shape.width > presentation.slide_width + .01 * EMU or
                    shape.top + shape.height > presentation.slide_height + .01 * EMU):
                results.append(Result("V4", "fail", tr("page", page=page) + " " + tr("v4_bad", shape=index)))
    return results or [Result("V4", "pass", tr("v4_ok"))]


def check_overflow(presentation):
    results = []
    for page, slide in enumerate(presentation.slides, 1):
        for index, shape in enumerate(slide.shapes, 1):
            frames = []
            if shape.has_text_frame and shape.text.strip():
                frames.append((shape.text_frame, shape.width, shape.height, str(index)))
            if shape.has_table:
                for row_index, row in enumerate(shape.table.rows):
                    for col_index, cell in enumerate(row.cells):
                        frames.append((cell.text_frame, shape.table.columns[col_index].width,
                                       row.height, f"{index}/{row_index + 1}/{col_index + 1}"))
            for frame, width, height, label in frames:
                size = max([run.font.size.pt for p in frame.paragraphs for run in p.runs if run.font.size] +
                           [p.font.size.pt for p in frame.paragraphs if p.font.size] + [10])
                usable_w = (width - frame.margin_left - frame.margin_right) / EMU
                usable_h = (height - frame.margin_top - frame.margin_bottom) / EMU
                needed = text_height(frame.text, usable_w, size)
                if needed > usable_h + size * 1.32 / 72:
                    results.append(Result("V9", "warn", tr("v9_bad", page=page, shape=label,
                                                           needed=needed, available=usable_h)))
    return results or [Result("V9", "pass", tr("v9_ok"))]


def check_overlaps(presentation):
    results = []
    for number, slide in enumerate(presentation.slides, 1):
        shapes = [(index, shape) for index, shape in enumerate(slide.shapes, 1) if get_role(shape) == "content"]
        for i, (first_index, first) in enumerate(shapes):
            for second_index, second in shapes[i + 1:]:
                overlap_w = min(first.left + first.width, second.left + second.width) - max(first.left, second.left)
                overlap_h = min(first.top + first.height, second.top + second.height) - max(first.top, second.top)
                if overlap_w > .025 * EMU and overlap_h > .025 * EMU:
                    results.append(Result("T4", "fail", tr("overlap_bad", page=number,
                                                           first=first_index, second=second_index)))
    return results or [Result("T4", "pass", tr("overlap_ok"))]


def check_outputs(course, engine_root):
    results = preflight(course, engine_root)
    if not course.pptx_path.is_file():
        results.append(Result("V1", "fail", tr("missing_output", path=course.pptx_path)))
        return results
    presentation = Presentation(course.pptx_path)
    actual, expected = len(presentation.slides), len(course.pages)
    results.append(Result("V1", "pass" if actual == expected else "fail", tr(
        "v1_ok" if actual == expected else "v1_bad", actual=actual, expected=expected)))
    notes_ok = actual == expected
    for index, slide in enumerate(presentation.slides):
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        if not notes.strip() or index >= expected or notes != course.pages[index].notes:
            notes_ok = False
            results.append(Result("V3", "fail", tr("page", page=index + 1) + " " + tr("v3_mismatch")))
    if notes_ok:
        results.append(Result("V3", "pass", tr("v3_ok", count=actual)))
    results.extend(check_geometry(presentation))
    backgrounds = [shape for shape in presentation.slides[0].shapes
                   if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and get_role(shape) == "background"]
    cover_ok = False
    for shape in backgrounds:
        width, height = shape.image.size
        full_bleed = (abs(shape.left) <= .01 * EMU and abs(shape.top) <= .01 * EMU and
                      abs(shape.width - presentation.slide_width) <= .01 * EMU and
                      abs(shape.height - presentation.slide_height) <= .01 * EMU)
        if full_bleed and width * height >= 2_070_000:
            cover_ok = True
            results.append(Result("V5", "pass", tr("v5_ok", width=width, height=height,
                                                  mp=width * height / 1_000_000)))
            break
    if not cover_ok:
        results.append(Result("V5", "fail", tr("v5_bad")))
    if not course.docx_path.is_file():
        results.append(Result("V6", "fail", tr("missing_output", path=course.docx_path)))
    else:
        units = read_document_units(course.docx_path)
        word_ok = len(units) == actual
        if not word_ok:
            results.append(Result("V6", "fail", tr("v6_bad", page="?")))
        for index, (unit, slide) in enumerate(zip(units, presentation.slides), 1):
            original_notes = slide.notes_slide.notes_text_frame.text
            if (unit["body"] != extract_slide_text(slide) or "\n".join(unit["notes"]) != original_notes or
                    unit["page_break"] != (index > 1) or unit["marker"] != f"deckflow_slide_{index:02}"):
                word_ok = False
                results.append(Result("V6", "fail", tr("v6_bad", page=index)))
        if word_ok:
            results.append(Result("V6", "pass", tr("v6_ok", count=actual)))
    results.extend(check_overflow(presentation))
    results.extend(check_overlaps(presentation))
    suspect = []
    for page in course.pages:
        text = "\n".join(value for _, value in page.entries) + page.notes
        if re.search(r"[\u3400-\u9fffA-Za-z]{1,12}\u8bf4\u8fc7", text):
            suspect.append(Result("V10", "warn", tr("v10_warn", page=page.number)))
        count = len(page.all("bullet"))
        if count > 5:
            results.append(Result("T4", "warn", tr("bullet_warn", page=page.number, count=count)))
    results.extend(suspect or [Result("V10", "pass", tr("v10_ok"))])
    if course.pages[-1].layout != "sources":
        results.append(Result("CONTENT", "warn", tr("source_last")))
    return results
