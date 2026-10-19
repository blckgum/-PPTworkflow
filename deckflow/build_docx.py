"""从成品 PPT 反读正文和备注，为每页生成一张 A4 讲稿。"""
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation

from .content import WorkflowError
from .rendering import extract_slide_text, text_height
from .strings import tr


def bookmark(paragraph, name, index):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(index))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(index))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def set_font(run, name, size, color=None, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    props = run._element.get_or_add_rPr()
    props.rFonts.set(qn("w:eastAsia"), name)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def document_size(body, notes, title, page):
    # A4 可用高约 10.2 英寸；预留页眉、标题、标签和空段间距。
    for size in (11, 10.5, 10, 9.5, 9):
        needed = text_height(title, 6.8, 16) + .75
        needed += sum(text_height(text, 6.8, size) + .035 for text in body)
        needed += sum(text_height(text or " ", 6.8, size) + .035 for text in notes.split("\n"))
        if needed <= 9.7:
            return size
    raise WorkflowError("V6", tr("document_overflow"), page, "notes")


def build_docx(course, theme):
    # 故意重新打开已经落盘的 PPT，避免维护第二份正文。
    presentation = Presentation(course.pptx_path)
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.2677), Inches(11.6929)
    section.top_margin = section.bottom_margin = Inches(.68)
    section.left_margin = section.right_margin = Inches(.72)
    section.header_distance = section.footer_distance = Inches(.25)
    normal = document.styles["Normal"]
    normal.font.name = theme["fonts"]["cjk"]
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.12
    header = section.header.paragraphs[0]
    set_font(header.add_run(course.meta["title"]), theme["fonts"]["cjk"], 8,
             theme["palette"]["ink_soft"])
    footer = section.footer.paragraphs[0]
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.core_properties.title = tr("doc_title", title=course.meta["title"])
    document.core_properties.author = course.meta["org"]
    index = 0

    def paragraph(text, size, mark=None, bold=False, color=None, page_break=False):
        nonlocal index
        item = document.add_paragraph()
        item.paragraph_format.page_break_before = page_break
        item.paragraph_format.keep_with_next = False
        item.paragraph_format.widow_control = False
        set_font(item.add_run(text), theme["fonts"]["cjk"], size, color, bold)
        if mark:
            index += 1
            bookmark(item, mark, index)
        return item

    for number, slide in enumerate(presentation.slides, 1):
        body = extract_slide_text(slide)
        notes = slide.notes_slide.notes_text_frame.text
        title = tr("doc_slide", number=number, title=course.pages[number - 1].get("title"))
        size = document_size(body, notes, title, number)
        paragraph(title, 16, f"deckflow_slide_{number:02}", bold=True,
                  color=theme["palette"]["accent"], page_break=number > 1)
        paragraph(tr("body_label"), 10, f"deckflow_body_{number:02}", bold=True)
        for text in body:
            paragraph(text, size)
        paragraph(tr("notes_label"), 10, f"deckflow_notes_{number:02}", bold=True)
        for line in notes.split("\n"):
            heading = line.startswith("\u3010")
            hint = line.startswith("\u3014")
            paragraph(line, size, bold=heading,
                      color=theme["palette"]["emphasis"] if hint else None)
    document.save(course.docx_path)
    return course.docx_path


def read_document_units(path):
    units = []
    current, mode = None, None
    for paragraph in Document(path).paragraphs:
        names = [node.get(qn("w:name"), "") for node in paragraph._p.findall(qn("w:bookmarkStart"))]
        if any(name.startswith("deckflow_slide_") for name in names):
            current = {"body": [], "notes": [], "page_break": bool(paragraph.paragraph_format.page_break_before),
                       "marker": next(name for name in names if name.startswith("deckflow_slide_"))}
            units.append(current)
            mode = None
        elif any(name.startswith("deckflow_body_") for name in names):
            mode = "body"
        elif any(name.startswith("deckflow_notes_") for name in names):
            mode = "notes"
        elif current is not None and mode:
            current[mode].append(paragraph.text)
    return units
