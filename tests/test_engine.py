"""解析、成品校验与回归测试；所有课程正文从课程文件读取。"""
from copy import deepcopy
from pathlib import Path
import shutil
import tempfile
import unittest

from docx import Document
from pptx import Presentation
from pptx.util import Inches

from deckflow.build_docx import build_docx
from deckflow.build_pptx import build_pptx
from deckflow.checks import check_geometry, check_outputs, check_overflow, preflight, passed
from deckflow.content import WorkflowError, load_course
from deckflow.layouts import LAYOUTS
from deckflow.rendering import line_count
from deckflow.source_lint import lint_sources
from deckflow.theme import get_theme

ROOT = Path(__file__).resolve().parents[1]


class EngineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name) / "course"
        shutil.copytree(ROOT / "courses" / "CQB", self.root,
                        ignore=shutil.ignore_patterns("\u8f93\u51fa"))

    def tearDown(self):
        self.temp.cleanup()

    def replace_deck(self, before, after):
        path = self.root / "deck.txt"
        path.write_text(path.read_text(encoding="utf-8").replace(before, after, 1), encoding="utf-8")

    def test_example_uses_every_layout(self):
        course = load_course(self.root)
        self.assertEqual({page.layout for page in course.pages}, set(LAYOUTS))
        self.assertEqual(sum(page.seconds for page in course.pages), 1530)

    def test_unknown_key_reports_page_and_key(self):
        self.replace_deck("layout\tpoints", "layout\tpoints\nbiullet\ttest")
        with self.assertRaises(WorkflowError) as context:
            load_course(self.root)
        self.assertEqual(context.exception.page, 4)
        self.assertEqual(context.exception.field, "biullet")

    def test_spaces_are_not_tabs(self):
        self.replace_deck("layout\tcover", "layout cover")
        with self.assertRaises(WorkflowError):
            load_course(self.root)

    def test_repeated_single_field_rejected(self):
        self.replace_deck("layout\tcover", "layout\tcover\nlayout\tcover")
        with self.assertRaises(WorkflowError):
            load_course(self.root)

    def test_page_numbers_must_be_sequential(self):
        self.replace_deck("[02]", "[03]")
        with self.assertRaises(WorkflowError):
            load_course(self.root)

    def test_layout_cannot_silently_drop_field(self):
        self.replace_deck("layout\tcover", "layout\tcover\nbullet\ttest")
        with self.assertRaises(WorkflowError) as context:
            load_course(self.root)
        self.assertEqual(context.exception.field, "bullet")

    def test_columns_need_two_or_three(self):
        course = load_course(self.root)
        page = next(p for p in course.pages if p.layout == "columns")
        value = page.all("column")[0]
        self.replace_deck("column\t" + value + "\n", "")
        value = page.all("column")[1]
        self.replace_deck("column\t" + value + "\n", "")
        with self.assertRaises(WorkflowError):
            load_course(self.root)

    def test_head_body_pairs_are_ordered(self):
        self.replace_deck("head\t", "body\t")
        with self.assertRaises(WorkflowError):
            load_course(self.root)

    def test_table_column_counts_match(self):
        course = load_course(self.root)
        row = next(p for p in course.pages if p.layout == "table").all("row")[1]
        self.replace_deck("row\t" + row, "row\tbroken")
        with self.assertRaises(WorkflowError):
            load_course(self.root)

    def test_duration_failure(self):
        self.replace_deck("time\t120", "time\t990")
        result = preflight(load_course(self.root), ROOT / "deckflow")
        self.assertTrue(any(r.code == "V2" and r.status == "fail" for r in result))

    def test_empty_notes_failure(self):
        (self.root / "notes" / "05.txt").write_text("\n", encoding="utf-8")
        result = preflight(load_course(self.root), ROOT / "deckflow")
        self.assertTrue(any(r.code == "V3" and r.status == "fail" for r in result))

    def test_missing_image_failure_has_expected_path(self):
        course = load_course(self.root)
        name = course.pages[0].images[0][0]
        path = course.image_path(name)
        path.unlink()
        result = preflight(course, ROOT / "deckflow")
        self.assertTrue(any(r.code == "V8" and r.status == "fail" and str(path) in r.message for r in result))

    def test_asset_traversal_rejected(self):
        course = load_course(self.root)
        with self.assertRaises(WorkflowError):
            course.image_path("../course.txt", 1)

    def test_cjk_wraps_character_by_character(self):
        self.assertEqual(line_count("\u4e2d" * 20, 100, 10), 2)
        self.assertEqual(line_count("\u4e2d\n\u6587", 100, 10), 2)

    def test_themes_are_independent_copies(self):
        theme = get_theme("military_light")
        theme["palette"]["accent"] = "000000"
        self.assertNotEqual(get_theme("military_light")["palette"]["accent"], "000000")

    def test_source_lint_catches_valid_and_invalid_chinese_quotes(self):
        target = Path(self.temp.name) / "lint"
        target.mkdir()
        file = target / "bad.py"
        file.write_text("bad = '\u4e2d\u6587\"\u6587\u672c'\n", encoding="utf-8")
        self.assertTrue(lint_sources(target))
        file.write_text('bad = "\u4eca\u5929"\u4e3a\u4ec0\u4e48"\u6253"\n', encoding="utf-8")
        self.assertTrue(lint_sources(target))
        file.write_text('bad = f"\u4e2d\u6587{1}"\n', encoding="utf-8")
        self.assertTrue(lint_sources(target))
        file.write_text('# \u6ce8\u91ca\u5141\u8bb8\u4e2d\u6587\nvalue = "ascii"\n', encoding="utf-8")
        self.assertEqual(lint_sources(target), [])


class OutputTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.root = Path(cls.temp.name) / "course"
        shutil.copytree(ROOT / "courses" / "CQB", cls.root,
                        ignore=shutil.ignore_patterns("\u8f93\u51fa"))
        cls.course = load_course(cls.root)
        theme = get_theme(cls.course.meta["theme"])
        build_pptx(cls.course, theme)
        build_docx(cls.course, theme)

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def test_example_passes_all_checks_without_warnings(self):
        results = check_outputs(self.course, ROOT / "deckflow")
        self.assertTrue(passed(results), results)
        self.assertTrue(all(result.status == "pass" for result in results), results)
        self.assertEqual({r.code for r in results}, {"V1", "V2", "V3", "V4", "V5", "V6", "V7", "V8", "V9", "V10", "T4"})

    def test_geometry_detects_outside_shape(self):
        presentation = Presentation(self.course.pptx_path)
        presentation.slides[0].shapes[0].left = Inches(-.1)
        self.assertFalse(passed(check_geometry(presentation)))

    def test_white_theme_hides_timing_and_preserves_duration_validation(self):
        root = Path(self.temp.name) / "white"
        shutil.copytree(ROOT / "courses" / "CQB", root,
                        ignore=shutil.ignore_patterns("\u8f93\u51fa"))
        course = load_course(root)
        course.meta["theme"] = "military_white"
        build_pptx(course, get_theme("military_white"))
        presentation = Presentation(course.pptx_path)
        from deckflow.strings import tr
        self.assertEqual(str(presentation.slides[3].background.fill.fore_color.rgb), "FFFFFF")
        for page, slide in zip(course.pages, presentation.slides):
            label = tr("time_label", minutes=page.seconds // 60, seconds=page.seconds % 60)
            self.assertFalse(any(shape.has_text_frame and label in shape.text for shape in slide.shapes))
        self.assertTrue(passed(preflight(course, ROOT / "deckflow")))
        course.pages[0].entries = [(key, "990" if key == "time" else value)
                                   for key, value in course.pages[0].entries]
        self.assertTrue(any(item.code == "V2" and item.status == "fail"
                            for item in preflight(course, ROOT / "deckflow")))

    def test_overflow_detects_cjk_text(self):
        presentation = Presentation(self.course.pptx_path)
        box = presentation.slides[1].shapes.add_textbox(0, 0, Inches(1), Inches(.2))
        box.text = "\u4e2d" * 300
        self.assertTrue(any(r.status == "warn" for r in check_overflow(presentation)))

    def test_word_tampering_is_detected(self):
        original = self.course.docx_path.read_bytes()
        try:
            document = Document(self.course.docx_path)
            document.paragraphs[3].text += " tampered"
            document.save(self.course.docx_path)
            result = check_outputs(self.course, ROOT / "deckflow")
            self.assertTrue(any(r.code == "V6" and r.status == "fail" for r in result))
        finally:
            self.course.docx_path.write_bytes(original)

    def test_ppt_notes_tampering_is_detected(self):
        original = self.course.pptx_path.read_bytes()
        try:
            presentation = Presentation(self.course.pptx_path)
            presentation.slides[4].notes_slide.notes_text_frame.text = "tampered"
            presentation.save(self.course.pptx_path)
            result = check_outputs(self.course, ROOT / "deckflow")
            self.assertTrue(any(r.code == "V3" and r.status == "fail" for r in result))
        finally:
            self.course.pptx_path.write_bytes(original)


if __name__ == "__main__":
    unittest.main()
