"""在隔离副本执行 T1–T4；--render 使用现成 LibreOffice 与 Poppler。"""
import argparse
import hashlib
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import xml.etree.ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


def code_hashes(root):
    return {path.relative_to(root).as_posix(): hashlib.sha256(path.read_bytes()).hexdigest()
            for path in root.rglob("*.py") if ".validation" not in path.parts}


def command(args, cwd, expected=0):
    result = subprocess.run(args, cwd=cwd, text=True, capture_output=True, timeout=180)
    if (expected == 0 and result.returncode != 0) or (expected != 0 and result.returncode == 0):
        raise AssertionError(f"{args}: exit {result.returncode}\n{result.stdout}\n{result.stderr}")
    return result.stdout + result.stderr


def normalize(text):
    return re.sub(r"\s+", "", unicodedata.normalize("NFC", text))


def render_artifacts(course, folder, log):
    from pptx import Presentation
    from deckflow.rendering import extract_slide_text, get_role
    for tool in ("soffice", "pdfinfo", "pdftotext", "pdftoppm"):
        if not shutil.which(tool):
            raise RuntimeError(f"Missing rendering tool: {tool}")
    folder.mkdir(parents=True, exist_ok=True)
    presentation = Presentation(course.pptx_path)
    for source, stem in ((course.pptx_path, "slides"), (course.docx_path, "notes")):
        filename = stem + source.suffix
        shutil.copy2(source, folder / filename)
        # 原生命令只收到 ASCII 文件名；中文原路径由 Python 处理。
        profile = (folder / ("profile-" + stem)).as_uri()
        output = command(["soffice", "-env:UserInstallation=" + profile, "--headless",
                          "--convert-to", "pdf", "--outdir", ".", filename], folder)
        log.append(output)
        pdf = stem + ".pdf"
        assert (folder / pdf).is_file(), pdf
        info = command(["pdfinfo", pdf], folder)
        match = re.search(r"^Pages:\s+(\d+)", info, re.M)
        assert match and int(match[1]) == len(presentation.slides), info
        for number, slide in enumerate(presentation.slides, 1):
            # 按绘制顺序提取，避免多栏文字按视觉行交错后产生假缺失。
            text = command(["pdftotext", "-f", str(number), "-l", str(number), "-raw", pdf, "-"], folder)
            expected = []
            if stem == "slides":
                for shape in slide.shapes:
                    if get_role(shape) != "content":
                        continue
                    if shape.has_text_frame:
                        expected.append(shape.text)
                    if shape.has_table:
                        expected.extend(cell.text for row in shape.table.rows for cell in row.cells)
            else:
                expected = extract_slide_text(slide) + slide.notes_slide.notes_text_frame.text.split("\n")
            for fragment in expected:
                assert normalize(fragment) in normalize(text), f"{stem} page {number}: clipped/missing {fragment!r}"
        # 以 PDF 页面矩形与实际文本坐标检查出血，不依赖 PNG 像素边缘。
        bbox = command(["pdftotext", "-bbox", pdf, "-"], folder)
        tree = ET.fromstring(bbox[bbox.index("<html"):])
        for page in tree.iter("{http://www.w3.org/1999/xhtml}page"):
            width, height = float(page.attrib["width"]), float(page.attrib["height"])
            for word in page.iter("{http://www.w3.org/1999/xhtml}word"):
                assert (float(word.attrib["xMin"]) >= -.75 and float(word.attrib["yMin"]) >= -.75 and
                        float(word.attrib["xMax"]) <= width + .75 and float(word.attrib["yMax"]) <= height + .75)
        log.append(f"T4 {stem}: {len(presentation.slides)} PDF pages; all content preserved; text inside page bounds")
    command(["pdftoppm", "-png", "-r", "60", "slides.pdf", "page"], folder)
    assert len(list(folder.glob("page-*.png"))) == len(presentation.slides)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--render", action="store_true")
    parser.add_argument("--report-dir", default=str(ROOT / ".validation" / "acceptance"))
    args = parser.parse_args()
    report = Path(args.report_dir).resolve()
    report.mkdir(parents=True, exist_ok=True)
    log = []
    with tempfile.TemporaryDirectory(prefix="deckflow-acceptance-") as directory:
        checkout = Path(directory) / "repo"
        shutil.copytree(ROOT, checkout, ignore=shutil.ignore_patterns(
            ".git", ".validation", "__pycache__", ".venv", "\u8f93\u51fa"))
        name = "\u6d4b\u8bd5\u8bfe"
        hashes = code_hashes(checkout)
        log.append(command([sys.executable, "run.py", "--new", name], checkout))
        course_root = checkout / "courses" / name
        assert all((course_root / item).exists() for item in ("course.txt", "deck.txt", "notes", "\u7d20\u6750"))
        example = checkout / "courses" / "CQB"
        # 仅编辑约定的内容文件、讲稿与三张素材；没有课程专属代码。
        for filename in ("course.txt", "deck.txt"):
            text = (example / filename).read_text(encoding="utf-8")
            if filename == "course.txt":
                text = re.sub(r"^title\t.*$", "title\t" + name, text, flags=re.M)
            (course_root / filename).write_text(text, encoding="utf-8")
        shutil.rmtree(course_root / "notes")
        shutil.copytree(example / "notes", course_root / "notes")
        for image in (example / "\u7d20\u6750").glob("*.png"):
            shutil.copy2(image, course_root / "\u7d20\u6750" / image.name)
        output = command([sys.executable, "run.py", name], checkout)
        assert "[\u5931\u8d25]" not in output and "[\u8b66\u544a]" not in output, output
        assert code_hashes(checkout) == hashes, "T1 modified Python files"
        log.extend([output, "T1 PASS: new course, 3 images, PPT+Word, all checks pass, Python hashes unchanged"])

        deck = course_root / "deck.txt"
        original = deck.read_bytes()
        # 已有总时长 1530 秒，给第一页加 870 秒，得到 40 分钟。
        deck.write_bytes(original.replace(b"time\t120", b"time\t990", 1))
        output = command([sys.executable, "run.py", name], checkout, expected=1)
        assert "[\u5931\u8d25] V2" in output, output
        log.append(output)
        deck.write_bytes(original)

        note = course_root / "notes" / "05.txt"
        saved = note.read_bytes()
        note.unlink()
        output = command([sys.executable, "run.py", name], checkout, expected=1)
        assert "[\u5931\u8d25] V3" in output and "05.txt" in output, output
        log.append(output)
        note.write_bytes(saved)

        from deckflow.content import load_course
        course = load_course(course_root)
        image = course.pages[0].images[0][0]
        deck.write_text(original.decode().replace(image, "missing.png", 1), encoding="utf-8")
        output = command([sys.executable, "run.py", name], checkout, expected=1)
        assert "[\u5931\u8d25] V8" in output and str(course.assets / "missing.png") in output, output
        log.append(output)
        deck.write_bytes(original)

        source = checkout / "deckflow" / "theme.py"
        theme_source = source.read_bytes()
        for bad in ('\nbad = "\u4eca\u5929"\u4e3a\u4ec0\u4e48"\u6253"\n',
                    "\nbad = '\u4e2d\u6587\"\u6587\u672c'\n"):
            source.write_bytes(theme_source + bad.encode())
            output = command([sys.executable, "run.py", name], checkout, expected=1)
            assert "[\u5931\u8d25] V7" in output and "theme.py:" in output, output
            log.append(output)
        source.write_bytes(theme_source)
        log.append("T2 PASS: V2/V3/V8/V7 reject deliberately broken inputs; V7 catches syntax and valid-string cases")

        layouts_hash = hashlib.sha256((checkout / "deckflow" / "layouts.py").read_bytes()).hexdigest()
        source.write_bytes(theme_source + b'\nTHEMES["acceptance"] = deepcopy(THEMES["military_light"])\n'
                           b'THEMES["acceptance"]["palette"].update(paper="EEF3FF", accent="244B80")\n')
        metadata = course_root / "course.txt"
        metadata_source = metadata.read_bytes()
        metadata.write_bytes(metadata_source.replace(b"military_light", b"acceptance"))
        output = command([sys.executable, "run.py", name], checkout)
        log.append(output)
        from pptx import Presentation
        presentation = Presentation(course.pptx_path)
        assert str(presentation.slides[3].background.fill.fore_color.rgb) == "EEF3FF"
        assert any(shape.has_text_frame and any(run.font.color.type is not None and
                   run.font.color.type == 1 and str(run.font.color.rgb) == "244B80"
                   for paragraph in shape.text_frame.paragraphs for run in paragraph.runs)
                   for slide in presentation.slides for shape in slide.shapes)
        assert hashlib.sha256((checkout / "deckflow" / "layouts.py").read_bytes()).hexdigest() == layouts_hash
        log.append("T3 PASS: copied theme, changed paper+accent, actual PPT colors verified, layouts unchanged")
        source.write_bytes(theme_source)
        metadata.write_bytes(metadata_source)
        log.append(command([sys.executable, "run.py", name], checkout))
        log.append(command([sys.executable, "run.py", name, "--check"], checkout))
        assert code_hashes(checkout) == hashes

        # 超过五条会警告；仍要先保证正文可排入页面。
        page = course.pages[3]
        first = page.all("bullet")[0]
        deck.write_text(original.decode().replace("bullet\t" + first,
                        ("bullet\t" + first + "\n") * 3 + "bullet\t" + first, 1), encoding="utf-8")
        output = command([sys.executable, "run.py", name], checkout)
        assert "[\u8b66\u544a] T4" in output and "6 \u4e2a\u8981\u70b9" in output, output
        log.append(output)
        deck.write_bytes(original)
        log.append(command([sys.executable, "run.py", name], checkout))
        course = load_course(course_root)
        for path in (course.pptx_path, course.docx_path):
            shutil.copy2(path, report / path.name)
        if args.render:
            render_artifacts(course, report / "render", log)
            log.append("T4 PASS: geometry, overlaps, overflow, >5 bullet warning, LibreOffice PDFs and all page PNGs")
        else:
            log.append("T4 PARTIAL: geometry and text checks passed; LibreOffice rendering not run (use --render)")
    (report / "results.txt").write_text("\n\n".join(log), encoding="utf-8")
    for line in log:
        if line.startswith(("T1 ", "T2 ", "T3 ", "T4 ")):
            print(line)
    print(report)


if __name__ == "__main__":
    main()
